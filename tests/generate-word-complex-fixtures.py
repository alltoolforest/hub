from pathlib import Path
import sys,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image,ImageDraw
out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
manifest={}
def base(title,description):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  d.styles[name].font.name='Liberation Sans';d.styles[name].font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.add_paragraph(title,'Title');d.add_paragraph(description)
 sec.header.paragraphs[0].text='AllToolForest Word engine acceptance fixtures'
 f=sec.footer.paragraphs[0];f.add_run('Page ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);f._p.append(field)
 return d
def save(d,name,features):
 for style in d.styles:
  for node in list(style.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 for node in list(d.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 d.save(out/name);manifest[name]={'features':features,'paragraphs':len(d.paragraphs),'tables':len(d.tables),'sections':len(d.sections)}
def border(table):
 pr=table._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
 pr.append(b)
 for cell in table.rows[0].cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DBEAFE');cell._tc.get_or_add_tcPr().append(sh)
  for run in cell.paragraphs[0].runs:run.bold=True
 repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
def link(p,text,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
# 1 formatting and literal search
D=base('Formatting and text editing','Use this document to check detection of inherited and direct run formatting, cursor edits, search, paragraph structure and saved output.')
s=D.styles.add_style('QA Base',WD_STYLE_TYPE.PARAGRAPH);s.base_style=D.styles['Normal'];s.font.name='Liberation Serif';s.font.size=Pt(13)
s=D.styles.add_style('QA Derived',WD_STYLE_TYPE.PARAGRAPH);s.base_style=D.styles['QA Base'];s.font.bold=True
p=D.add_paragraph('Inherited style sentinel',style='QA Derived')
p=D.add_paragraph();
for text,font,size,b,i,color in [('Alpha ','Liberation Serif',12,True,False,'943126'),('Beta ','Liberation Sans',14,False,True,'145A32'),('Gamma','Liberation Mono',10.5,False,False,'154360')]:
 r=p.add_run(text);r.font.name=font;r.font.size=Pt(size);r.bold=b;r.italic=i;r.font.color.rgb=RGBColor.from_string(color)
D.add_paragraph('Cursor split sentinel one two three four five.')
D.add_paragraph('Search alpha ALPHA alphabet alpha. Literal a+b $5 [brackets] and café CAFÉ.')
D.add_paragraph('Duplicate repeat repeat repeat words verify exact range replacement.')
D.add_paragraph('')
p=D.add_paragraph();r=p.add_run('Underline ');r.underline=True;r=p.add_run('strike ');r.font.strike=True;r=p.add_run('superscript');r.font.superscript=True
D.add_heading('Second page editing targets',1).paragraph_format.page_break_before=True
D.add_paragraph('Paragraph A remains separate.');D.add_paragraph('Paragraph B is the adjacent merge target.')
p=D.add_paragraph('Tabs and breaks');p.add_run('\tTab value\nManual line break')
for i in range(4):D.add_paragraph(f'Formatting record {i+1}: replace a short sentence with a longer sentence, then restore it and compare all surrounding runs.')
save(D,'01_Formatting_and_Search.docx',['style ancestry','mixed runs','decimal font','range edits','split merge','search case whole word literal','tabs breaks','undo redo'])
# 2 section/page flow
D=base('Page flow and section layout','Use this document to compare growth and shortening across page boundaries, paragraph keep rules, section geometry, headers and footers.')
for i in range(11):
 p=D.add_paragraph(f'Flow record {i+1:02d}. '+('Content should wrap naturally and preserve every word as paragraphs grow or shrink. '*3));p.paragraph_format.widow_control=True
 if i==3:p.paragraph_format.keep_with_next=True
sec=D.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
D.add_heading('Landscape section',1);D.add_paragraph('Landscape sentinel with a different page size and margins. The editor must not treat every section as portrait.')
for i in range(4):D.add_paragraph(f'Landscape record {i+1}. '+('This section exercises independent page geometry. '*4))
save(D,'02_Page_Flow_and_Sections.docx',['growth shrink','cross-page flow','landscape section','keep next','widow orphan','headers footers page fields'])
# 3 tables
D=base('Tables and structured records','Use this document to check ordinary cell edits while preserving merged cells, nested tables, repeated headers and surrounding paragraphs.')
T=D.add_table(rows=1,cols=3);T.rows[0].cells[0].text='Record';T.rows[0].cells[1].text='Description';T.rows[0].cells[2].text='Value'
for i in range(7):
 row=T.add_row();row.cells[0].text=f'R{i+1:02d}';row.cells[1].text='Cell edit sentinel' if i==0 else 'Detailed item with wrapping text and preserved borders.';row.cells[2].text=str((i+1)*125)
border(T);T.rows[3].cells[0].merge(T.rows[3].cells[1]).text='Merged region sentinel'
D.add_heading('Nested table sample',1);T=D.add_table(rows=2,cols=2);T.cell(0,0).text='Outer record';T.cell(0,1).text='Nested detail';T.cell(1,0).text='Parent cell';N=T.cell(1,1).add_table(rows=2,cols=2)
for i,row in enumerate(N.rows):
 for j,cell in enumerate(row.cells):cell.text=f'Nested {i+1} {j+1}'
border(T);border(N);D.add_paragraph('Trailing table paragraph sentinel remains after all table rows.')
save(D,'03_Tables_and_Merged_Cells.docx',['table cell edits','merged cells','nested table','borders','repeat headers','table preservation'])
# 4 objects review
D=base('Objects links and review content','Use this document to check object preservation and identify content that must be protected from ordinary text editing.')
p=D.add_paragraph('Editable object-adjacent sentinel. ');link(p,'External documentation','https://example.com/documentation')
img=Image.new('RGB',(900,240),'#e7eee8');draw=ImageDraw.Draw(img)
for x,c in [(70,'#166e52'),(350,'#477c69'),(630,'#8fbcaa')]:draw.rectangle((x,40,x+180,200),fill=c)
img.save(out/'fixture_image.png');D.add_picture(str(out/'fixture_image.png'),width=Inches(5.6));D.add_paragraph('Figure one with three colour blocks','Caption')
p=D.add_paragraph('Bookmark sentinel');bm=OxmlElement('w:bookmarkStart');bm.set(qn('w:id'),'12');bm.set(qn('w:name'),'qa_bookmark');p._p.insert(0,bm);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'12');p._p.append(end)
p=D.add_paragraph('Tracked replacement ')
for tag,text in [('del','old wording'),('ins','new wording')]:
 el=OxmlElement('w:'+tag);el.set(qn('w:id'),'20' if tag=='del' else '21');el.set(qn('w:author'),'QA Fixture');el.set(qn('w:date'),'2026-10-21T00:00:00Z');r=OxmlElement('w:r');t=OxmlElement('w:delText' if tag=='del' else 'w:t');t.text=text;r.append(t);el.append(r);p._p.append(el)
p=D.add_paragraph('Content control ');sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');tag=OxmlElement('w:tag');tag.set(qn('w:val'),'QA_CONTROL');pr.append(tag);content=OxmlElement('w:sdtContent');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Protected control value';r.append(t);content.append(r);sdt.append(pr);sdt.append(content);p._p.append(sdt)
p=D.add_paragraph('Field result ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'DATE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='21 October 2026';r.append(t);f.append(r);p._p.append(f)
D.add_paragraph('Independent editable sentinel for package preservation comparison.')
save(D,'04_Objects_Links_and_Review.docx',['image','hyperlink','bookmark','tracked revisions','content control','simple field','package preservation'])
# 5 unicode and sustained edits
D=base('Multilingual text and recovery','Use this document to check Unicode text, combining characters, long paragraphs, repeated search targets and recovery of an extended editing session.')
for label,text,font in [('Latin','café naïve Straße Ελληνικά Кириллица','DejaVu Sans'),('Arabic','مرحبا بالعالم هذا نص للتحرير','DejaVu Sans'),('Hindi','नमस्ते दुनिया यह संपादन परीक्षण है','Noto Sans Devanagari'),('Hebrew','שלום עולם בדיקת עריכה','DejaVu Sans')]:
 p=D.add_paragraph(label+' ');r=p.add_run(text);r.font.name=font
p=D.add_paragraph('Combining sentinel e\u0301 a\u0308 versus é ä. Supplementary character: ');r=p.add_run('😀');r.font.name='DejaVu Sans'
p=D.add_paragraph();r=p.add_run('Unavailable font sentinel with readable fallback');r.font.name='ATF Missing Font 2026'
for i in range(14):D.add_paragraph(f'Recovery record {i+1:02d}. '+('Persistent edits must survive undo, redo, export, reload and repeated replacement without changing unrelated Unicode text. '*2))
save(D,'05_Multilingual_and_Recovery.docx',['Unicode','RTL','Indic scripts','combining marks','supplementary characters','long text','autosave recovery'])
(out/'fixture_manifest.json').write_text(json.dumps(manifest,indent=2))
print(json.dumps(manifest,indent=2))
